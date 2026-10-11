"""Excel export of the filtered maintenance plan summary and its occurrences."""
from io import BytesIO
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def export_plans(data, filters, plan_order=None):
    plans = {p['plan']: p for p in data['by_plan']}
    if plan_order is None:
        plan_order = list(plans)
    if (not isinstance(plan_order, list) or any(not isinstance(p, str) for p in plan_order)
            or len(plan_order) != len(plans) or set(plan_order) != set(plans)):
        raise ValueError('A lista de planos mudou. Atualize a tabela antes de exportar.')
    wb = Workbook()
    wb.remove(wb.active)
    scope = ' | '.join(f'{k}: {", ".join(map(str, v)) if isinstance(v, list) else v}'
                       for k, v in filters.items() if v not in ('', None, [])) or 'Todos os dados'

    row_numbers = {}

    def append(ws, values):
        ws.append(values)
        row_numbers[ws.title] = row_numbers.get(ws.title, 0) + 1
        # SAP descriptions and codes are literal text, never Excel formulas.
        for col, value in enumerate(values, 1):
            if isinstance(value, str):
                ws.cell(row_numbers[ws.title], col).data_type = 's'

    def sheet(name, headers, note):
        ws = wb.create_sheet(name)
        append(ws, [name])
        append(ws, [scope])
        append(ws, [note])
        append(ws, [])
        append(ws, headers)
        for row in (1, 2, 3):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        ws.row_dimensions[2].height = 42
        ws.row_dimensions[3].height = 32
        ws.freeze_panes = 'C6'
        return ws

    traits = [('Itens', 'item'), ('Locais de instalação', 'asset_original'),
              ('Ativos', 'asset_name'), ('Tipo de ordem', 'type'),
              ('Tipo de atividade', 'activity_type'), ('Centro de trabalho', 'work_center'),
              ('Código ABC', 'abc'), ('Condição da linha', 'condition'),
              ('Ciclo', 'cycle'), ('Unidade do ciclo', 'unit')]
    summary = sheet('Realização por plano', ['Plano', 'Atividade(s)', 'Previstas', 'Realizadas',
                    'Não realizadas', 'Índice de realização'] + [t[0] for t in traits],
                    'Mesmo recorte da tabela. Características distintas separadas por quebra de linha; associações exatas na segunda aba.')
    grouped = {}
    for r in data['rows']:
        grouped.setdefault(r['plan'], []).append(r)
    for code in plan_order:
        p = plans[code]
        values = [code, '\n'.join(p['descriptions']), p['planned'], p['done'], p['pending'],
                  p['rate'] / 100 if p['rate'] is not None else None]
        values.extend('\n'.join(sorted({str(r.get(key, '')) for r in grouped.get(code, [])
                                         if r.get(key) not in ('', None)})) for _, key in traits)
        append(summary, values)
    t = data['totals']
    summary_end = summary.max_row
    append(summary, ['Total', f'{len(plans)} planos', t['planned'], t['done'], t['pending'],
                     t['rate'] / 100 if t['rate'] is not None else None])
    for row in summary.iter_rows(min_row=6):
        row[5].number_format = '0.0%'

    fields = [('Plano', 'plan'), ('Item', 'item'), ('Atividade do plano', 'description'),
              ('Local do item', 'asset_original'), ('Ativo consolidado', 'asset'),
              ('Descrição do ativo', 'asset_name'), ('Ordem', 'order'),
              ('Descrição da ordem', 'order_description'), ('Data prevista', 'planned'),
              ('Origem da data', 'date_source'), ('Situação', 'execution'),
              ('Status usuário', 'user_status'), ('Status sistema', 'system_status'),
              ('Fim real', 'actual_end'), ('H duração', 'duration_hours'),
              ('Tipo de ordem · IP18', 'type'), ('Tipo de atividade · IP18', 'activity_type'),
              ('Centro de trabalho · IP18', 'work_center'), ('ABC · IP18', 'abc'),
              ('Condição · IP18', 'condition'), ('Ciclo', 'cycle'), ('Unidade do ciclo', 'unit'),
              ('Local da ordem · IW39', 'order_asset'), ('Tipo de ordem · IW39', 'order_type'),
              ('Tipo de atividade · IW39', 'order_activity_type'),
              ('Centro de trabalho · IW39', 'order_work_center'),
              ('Início básico', 'order_basic_start'), ('Fim básico', 'order_basic_end'),
              ('Fim-base (hora)', 'order_base_end_time'), ('Hora fim', 'order_scheduled_end_time'),
              ('Fonte IP18', 'source'), ('Fonte IW39', 'order_source'),
              ('Fonte programação', 'schedule_source'), ('ID da ocorrência', 'id')]
    detail = sheet('Ordens e ocorrências', [f[0] for f in fields],
                   'Uma linha por ocorrência dos planos listados, no mesmo recorte. Ocorrências sem ordem são preservadas; itens sem data não entram neste resumo.')
    dates = {'planned', 'actual_end', 'order_basic_start', 'order_basic_end'}
    for code in plan_order:
        for r in grouped.get(code, []):
            append(detail, [date.fromisoformat(r[key]) if key in dates and r.get(key)
                            else r.get(key) for _, key in fields])
    for ws, end in ((summary, summary_end), (detail, detail.max_row)):
        ws.auto_filter.ref = f'A5:{get_column_letter(ws.max_column)}{max(5,end)}'
        ws.sheet_view.showGridLines = False
        for row in ws:
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if isinstance(cell.value, date): cell.number_format = 'dd/mm/yyyy'
                if cell.row == 5:
                    cell.fill = PatternFill('solid', fgColor='126F66')
                    cell.font = Font(color='FFFFFF', bold=True)
                elif cell.row > 5 and cell.row % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='F0F6F5')
        ws['A1'].font = Font(size=16, bold=True, color='126F66')
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 24
        ws.column_dimensions['A'].width = 20
    summary.column_dimensions['B'].width = 65
    detail.column_dimensions['C'].width = 55
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='DCEEE9')
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
