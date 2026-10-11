import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from app import create_app
from database import connect
from sap_plan import DEFAULT_CONFIG, execution, import_exports, report, duration_hours, totals


class SAPPlanTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.path = self.root / 'test.sqlite3'
        self.app = create_app(self.path, self.root)
        self.client = self.app.test_client()
        self.client.get('/plano-manutencao')
        with self.client.session_transaction() as s: self.csrf = s['csrf']
        with connect(self.path) as db:
            base='G-CSN-79-001-010-'
            fixtures=[
                ('IH06','A',{'asset':base+'001','name':'Forno'}),
                ('IH06','B',{'asset':base+'002','name':'Bomba'}),
                ('IW39','1',{'order':'1','item':'I','plan':'P','system_status':'CONF','basic_start':'2026-09-02','source':'IW39'}),
                ('IW39','2',{'order':'2','item':'II','plan':'P','system_status':'ENTE','user_status':'EXCT','basic_start':'2026-09-10','source':'IW39'}),
                ('IW39','3',{'order':'3','plan':'EXTRA','basic_start':'2026-09-10','system_status':'CONF','source':'IW39'}),
                ('IP24','1',{'order':'1','item':'I','planned':'2026-09-01','source':'IP24'}),
                ('IP24','3',{'order':'','item':'III','planned':'2026-09-15','source':'IP24'}),
                ('IP24','4',{'order':'','item':'III','planned':'2026-10-15','source':'IP24'})]
            for i,asset in [('I','001-001'),('II','001-002'),('III','002-001'),('IV','002-002')]:
                fixtures.append(('IP18',i,{'item':i,'plan':'P','asset':base+asset,'item_text':'Lubrificar '+i,
                    'type':'ZPRV','activity_type':'C6','work_center':'MM','abc':'A' if i=='I' else 'B',
                    'condition':'OPE - Em operação','source':'IP18'}))
            for kind,key,row in fixtures:
                db.execute('INSERT INTO sap_rows VALUES(?,?,?)',(kind,key,json.dumps(row)))

    def tearDown(self): self.temp.cleanup()

    def post(self, path, data):
        return self.client.post('/api/sap/'+path, json=data, headers={'X-CSRF-Token':self.csrf})

    def test_plan_detail(self):
        response = self.client.get('/api/sap/plan-detail?plan=p')
        self.assertEqual(response.status_code, 200)
        data = response.json
        self.assertEqual(data['plan'], 'P')
        self.assertEqual(len(data['items']), 4)
        self.assertEqual(len(data['rows']), 4)
        self.assertEqual(len(data['undated']), 1)
        self.assertTrue(all(r['plan'] == 'P' for r in data['rows']))
        self.assertEqual(data['totals']['done'], 1)
        self.assertEqual(self.client.get('/api/sap/plan-detail?plan=EXTRA').status_code, 404)
        self.assertEqual(self.client.get('/api/sap/plan-detail').status_code, 400)

    def test_plan_detail_respects_period_and_assets(self):
        filters = {'plan': 'P', 'date_from': '2026-09-01', 'date_to': '2026-09-15'}
        data = self.post('plan-detail', filters).json
        self.assertEqual(data['totals']['planned'], 3)
        self.assertTrue(all(filters['date_from'] <= r['planned'] <= filters['date_to'] for r in data['rows']))
        expected = self.post('report', filters).json
        self.assertEqual(data['totals'], expected['totals'])
        data = self.post('plan-detail', dict(filters, assets=['G-CSN-79-001-010-001'])).json
        self.assertEqual(data['totals']['planned'], 2)
        self.assertTrue(all(r['asset'] == 'G-CSN-79-001-010-001' for r in data['rows']))
        self.assertEqual(self.post('plan-detail', dict(filters, date_to='2026-08-31')).status_code, 400)

    def test_summary_keeps_original_locations_above_xxx(self):
        with connect(self.path) as db:
            for i, location in enumerate(['G-CSN-79-001-010', 'G-CSN-79-001']):
                for kind, row in [('IP18', {'item': 'TOP'+str(i), 'plan': 'TOP', 'asset': location}),
                                  ('IW39', {'order': 'TOP'+str(i), 'item': 'TOP'+str(i), 'plan': 'TOP', 'basic_start': '2026-09-01'})]:
                    db.execute('INSERT INTO sap_rows VALUES(?,?,?)', (kind, 'TOP'+str(i), json.dumps(row)))
            data = report(db, {'plan': 'TOP'})
        self.assertEqual({r['asset'] for r in data['by_asset']}, {'G-CSN-79-001-010', 'G-CSN-79-001'})
        self.assertTrue(all(r['direct_location'] for r in data['by_asset']))
        self.assertEqual(data['totals']['planned'], 2)

    def test_parent_location_filter_and_group_include_descendants(self):
        parent = 'G-CSN-79-001-010'
        with connect(self.path) as db:
            for kind, row in [('IP18', {'item':'TOP', 'plan':'TOP', 'asset':parent}),
                              ('IW39', {'order':'TOP', 'item':'TOP', 'plan':'TOP', 'basic_start':'2026-09-01'})]:
                db.execute('INSERT INTO sap_rows VALUES(?,?,?)', (kind, 'TOP', json.dumps(row)))
        self.assertIn(parent, [r['asset'] for r in self.client.get('/api/sap/meta').json['assets']])
        filters = {'assets': [parent], 'date_from':'2026-09-01', 'date_to':'2026-09-30'}
        data = self.post('report', filters).json
        self.assertEqual(data['totals']['planned'], 4)
        direct = self.post('report', dict(filters, summary_asset=parent)).json
        self.assertEqual(direct['totals']['planned'], 1)
        self.assertEqual([p['plan'] for p in direct['by_plan']], ['TOP'])
        child = self.post('report', dict(filters, summary_asset=parent+'-001')).json
        self.assertEqual(child['totals']['planned'], 2)
        detail = self.post('plan-detail', dict(filters, summary_asset=parent+'-001', plan='P')).json
        self.assertEqual(detail['totals'], child['totals'])

        self.assertEqual(data['undated_count'], 1)
        self.assertEqual(self.post('report', dict(filters, assets=[parent, parent+'-001'])).json['totals'], data['totals'])
        self.assertEqual(self.post('groups', {'name':'Linha', 'assets':[parent]}).status_code, 200)
        group = self.client.get('/api/sap/meta').json['groups'][0]['id']
        self.assertEqual(self.post('report', dict(filters, assets=[], group=group)).json['totals'], data['totals'])

    def test_all_assets_card_ignores_asset_restrictions(self):
        period = {'date_from':'2026-09-01', 'date_to':'2026-09-30'}
        expected = self.post('report', period).json['totals']
        self.post('groups', {'name':'Forno', 'assets':['G-CSN-79-001-010-001']})
        group = self.client.get('/api/sap/meta').json['groups'][0]['id']
        narrowed = dict(period, assets=['G-CSN-79-001-010-001'], group=group, search='Forno')
        data = self.post('report', narrowed).json
        self.assertEqual(data['all_assets_totals'], expected)
        self.assertEqual(data['totals']['planned'], 2)
        self.assertEqual(expected['planned'], 3)
        filtered = self.post('report', dict(narrowed, abc='B')).json
        self.assertEqual(filtered['all_assets_totals']['planned'], 2)
        empty = self.post('report', dict(narrowed, date_from='2027-01-01', date_to='2027-01-31')).json
        self.assertEqual(empty['all_assets_totals']['planned'], 0)

    def test_duration_hours_and_deduplication(self):
        self.assertEqual(duration_hours('24:00:00', '08:00:00'), 8)
        self.assertEqual(duration_hours('08:00:00', '00:00:00'), 8)
        self.assertEqual(duration_hours('08:30:00', '09:00:00'), .5)
        self.assertEqual(duration_hours('00:00:00', '24:00:00'), 0)
        self.assertIsNone(duration_hours('', '08:00:00'))
        self.assertIsNone(duration_hours('25:00:00', '08:00:00'))
        row = {'order':'1', 'execution':'Executado', 'duration_hours':8}
        result = totals([row, row, {'order':'2','execution':'Não executado'}])
        self.assertEqual(result['duration_hours'], 8)
        self.assertEqual(result['duration_orders'], 1)
        self.assertEqual(result['duration_missing_orders'], 1)
        self.assertEqual(result['duration_done'], 8)
        self.assertEqual(result['duration_pending'], 0)
        pending = {'order':'3', 'execution':'Não executado', 'duration_hours':2.5}
        combined = totals([row, row, pending, pending])
        self.assertEqual(combined['duration_hours'], 10.5)
        self.assertEqual(combined['duration_done'], 8)
        self.assertEqual(combined['duration_pending'], 2.5)
        self.assertIsNone(totals([])['duration_done'])


    def test_manage_groups(self):
        a, b = 'G-CSN-79-001-010-001', 'G-CSN-79-001-010-002'
        self.post('groups', {'name':'Primeiro', 'assets':[a]})
        self.post('groups', {'name':'Segundo', 'assets':[b]})
        groups = self.client.get('/api/sap/meta').json['groups']
        first, second = [g['id'] for g in groups]
        self.assertEqual(self.post('groups/reorder', {'ids':[second, first]}).status_code, 200)
        self.assertEqual([g['id'] for g in self.client.get('/api/sap/meta').json['groups']], [second, first])
        self.assertEqual([g['id'] for g in self.post('report', {}).json['groups']], [second, first])
        self.assertEqual(self.post('groups/reorder', {'ids':[first, first]}).status_code, 400)
        self.assertEqual(self.post('groups', {'id':first, 'name':'Renomeado', 'assets':[a,b]}).status_code, 200)
        self.assertEqual(self.post('groups', {'id':first, 'name':'Renomeado', 'assets':[b]}).status_code, 200)
        edited = next(g for g in self.client.get('/api/sap/meta').json['groups'] if g['id']==first)
        self.assertEqual(edited['assets'], [b])
        self.assertEqual(edited['name'], 'Renomeado')
        self.assertEqual(self.post('groups', {'id':first, 'name':'Segundo', 'assets':[a]}).status_code, 400)
        self.assertEqual(self.post('groups', {'id':first, 'name':'Renomeado', 'assets':[]}).status_code, 400)
        before = self.post('report', {}).json['totals']
        self.assertEqual(self.post('groups/delete', {'id':first}).status_code, 200)
        self.assertEqual([g['id'] for g in self.client.get('/api/sap/meta').json['groups']], [second])
        self.assertEqual(self.post('report', {}).json['totals'], before)
        self.assertEqual(self.post('groups/delete', {'id':first}).status_code, 400)

    def test_export_plans_filtered_workbook(self):
        from io import BytesIO
        import openpyxl
        filters = {'date_from':'2026-09-01', 'date_to':'2026-09-30',
                   'summary_asset':'G-CSN-79-001-010-001'}
        response = self.post('export-plans', {'filters':filters, 'plan_order':['P']})
        self.assertEqual(response.status_code, 200)
        wb = openpyxl.load_workbook(BytesIO(response.data))
        self.assertEqual(wb.sheetnames, ['Realização por plano', 'Ordens e ocorrências'])
        summary, detail = wb.worksheets
        self.assertEqual([summary.cell(6,i).value for i in range(1,7)],
                         ['P','Lubrificar I\nLubrificar II',2,1,1,.5])
        rows = list(detail.iter_rows(min_row=6, values_only=True))
        self.assertEqual(len(rows), 2)
        self.assertEqual({r[6] for r in rows}, {'1','2'})
        self.assertEqual({r[1] for r in rows}, {'I','II'})
        self.assertTrue(all(r[0]=='P' and r[8].month==9 for r in rows))
        self.assertEqual(summary.freeze_panes, 'C6')
        self.assertEqual(summary['F6'].number_format, '0.0%')
        wb.close()
        response = self.post('export-plans', {'filters':dict(filters,summary_asset='G-CSN-79-001-010-002')})
        wb = openpyxl.load_workbook(BytesIO(response.data))
        self.assertEqual(wb.worksheets[1].max_row, 6)
        self.assertIsNone(wb.worksheets[1]['G6'].value)
        wb.close()
        self.assertEqual(self.post('export-plans', {'filters':filters,'plan_order':['EXTRA']}).status_code,400)

    def test_export_keeps_sap_text_literal(self):
        from sap_export import export_plans
        from io import BytesIO
        import openpyxl
        with connect(self.path) as db:
            data = report(db, {}, all_rows=True)
        data['by_plan'][0]['descriptions'] = ['=1+1']
        data['rows'][0]['order_description'] = '=HYPERLINK("x")'
        wb = openpyxl.load_workbook(export_plans(data, {}))
        self.assertEqual(wb.worksheets[0]['B6'].data_type, 's')
        self.assertEqual(wb.worksheets[0]['B6'].value, '=1+1')
        self.assertEqual(wb.worksheets[1]['H6'].data_type, 's')
        wb.close()

    def test_source_choices_saved_and_explicit_import(self):
        from unittest.mock import Mock
        path = self.root / 'nome livre.xlsx'
        path.touch()
        choices = dict(IP18=str(path), IW39='', IP24='', IH06='')
        def immediate_thread(target, **kwargs):
            return Mock(start=target)
        with patch('sap_plan.threading.Thread', side_effect=immediate_thread), patch('sap_plan.import_exports', return_value={'files':1}) as importer:
            response = self.post('import', {'choices':choices})
            self.assertEqual(response.status_code, 202)
            self.assertEqual(importer.call_args.args[2], [(path.resolve(), 'IP18')])
        saved = self.client.get('/api/sap/source-choices').json
        self.assertEqual(saved['choices'], choices)
        self.assertIn(str(path.resolve()), saved['available'])
        second_app = create_app(self.path, self.root)
        self.assertEqual(second_app.test_client().get('/api/sap/source-choices').json['choices'], choices)
        self.assertEqual(self.post('import', {'choices':dict(choices, IW39=str(path))}).status_code, 400)
        self.assertEqual(self.post('import', {'choices':dict(choices, IP18='missing.xlsx')}).status_code, 400)
        self.assertEqual(self.client.get('/api/sap/source-choices').json['choices'], choices)
        row = {'item':'NEW', 'plan':'NEW', 'asset':'G-CSN-79-001-010-001'}
        with patch('sap_plan.read_export', return_value=iter([row])) as reader:
            import_exports(self.path, self.root, [(path, 'IP18')])
            reader.assert_called_once_with(path, 'IP18')
        with connect(self.path) as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM sap_rows WHERE kind='IW39'").fetchone()[0], 3)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM sap_rows WHERE kind='IP18'").fetchone()[0], 1)

    def test_group_visibility_is_persistent_and_preserves_group(self):
        self.post('groups', {'name':'Teste', 'assets':['G-CSN-79-001-010-001']})
        g = self.client.get('/api/sap/meta').json['groups'][0]
        self.assertEqual(g['visible'], 1)
        self.assertEqual(self.post('groups/visibility', {'id':g['id'], 'visible':False}).status_code, 200)
        self.assertFalse(self.post('report', {}).json['groups'][0]['visible'])
        saved = self.client.get('/api/sap/meta').json['groups'][0]
        self.assertEqual(saved['assets'], g['assets'])
        self.assertEqual(saved['visible'], 0)
        self.assertEqual(self.post('report', {'group':g['id']}).json['totals']['planned'], 2)
        self.assertEqual(self.post('groups/visibility', {'id':g['id'], 'visible':True}).status_code, 200)
        self.assertTrue(self.post('report', {}).json['groups'][0]['visible'])
        self.assertEqual(self.post('groups/visibility', {'id':g['id'], 'visible':'false'}).status_code, 400)

    def test_master_scope_dates_and_missing_orders(self):
        data=self.post('report',{'date_from':'2026-09-01','date_to':'2026-09-30'}).json
        self.assertEqual(data['totals']['planned'],3)
        self.assertEqual(data['monthly'][0]['order_count'], 2)
        self.assertEqual(data['totals']['done'],1)
        self.assertEqual(data['totals']['pending'],2)
        self.assertEqual(data['totals']['rate'],33.3)
        self.assertEqual(data['total'],3)
        self.assertEqual(data['undated_count'],1)
        self.assertEqual(data['outside_master_orders'],1)
        self.assertEqual(data['daily'][0]['day'],'2026-09-01')
        self.assertEqual(data['daily'][1]['day'],'2026-09-10')
        self.assertIsNone(self.post('report',{'date_from':'2027-01-01'}).json['totals']['rate'])

    def test_group_level_and_index(self):
        prefix='G-CSN-79-001-010-'
        assets=self.client.get('/api/sap/meta').json['assets']
        self.assertEqual([a['asset'] for a in assets],[prefix+'001',prefix+'002'])
        self.assertEqual(self.post('groups',{'name':'Forno','assets':[prefix+'001-001']}).status_code,400)
        self.assertEqual(self.post('groups',{'name':'Forno','assets':[prefix+'001']}).status_code,200)
        group=self.client.get('/api/sap/meta').json['groups'][0]['id']
        data=self.post('report',{'group':group}).json
        self.assertEqual(data['totals']['planned'],2)
        self.assertEqual(data['totals']['rate'],50)
        self.assertEqual(data['groups'][0]['rate'],50)
        self.assertEqual(self.post('report',{'group':group,'assets':[prefix+'002']}).json['total'],0)

    def test_binary_execution_rules(self):
        cfg=DEFAULT_CONFIG|{'done_user':['EXCT'],'not_done_user':['NAEX']}
        self.assertEqual(self.post('config',cfg).status_code,200)
        data=self.post('report',{'date_to':'2026-09-30'}).json
        self.assertEqual(data['totals']['done'],2)
        for row in [{},{'system_status':'ENTE'},{'system_status':'CNPA'},{'user_status':'EXCT NAEX'}, {'system_status':'CONF','deleted':'X'}]:
            self.assertEqual(execution(row,cfg),'Não executado')
        self.assertEqual(self.post('config',cfg|{'not_done_user':['EXCT']}).status_code,400)

    def test_calendar_and_status_drilldown_keep_denominator(self):
        data=self.post('report',{'execution':'Executado','day':'2026-09-01'}).json
        self.assertEqual(data['total'],1)
        self.assertEqual(data['totals']['planned'],4)
        self.assertEqual(data['totals']['done'],1)
        self.assertEqual(len(data['daily']),4)
        self.assertEqual(len(data['by_plan']),1)
        self.assertEqual(data['by_plan'][0]['plan'],'P')
        self.assertEqual(data['by_plan'][0]['planned'],4)
        self.assertEqual(data['by_plan'][0]['done'],1)
        self.assertEqual(data['by_plan'][0]['pending'],3)
        self.assertEqual(data['by_plan'][0]['rate'],25)

    def test_all_master_filters_and_csrf(self):
        self.assertEqual(self.client.post('/api/sap/config',json={}).status_code,403)
        self.assertEqual(self.post('report',{'date_from':'2026-10-01','date_to':'2026-09-01'}).status_code,400)
        self.assertEqual(self.post('groups',{'name':'x','assets':['unknown']}).status_code,400)
        f={'type':'ZPRV','activity_type':'C6','work_center':'MM','abc':'A','text':'operacao','search':'forno'}
        self.assertEqual(self.post('report',f).json['totals']['planned'],1)
        self.assertEqual(self.post('report',f|{'activity_type':'invalid'}).json['total'],0)
        self.assertEqual(self.post('report',{'condition':'OPE - Em operação'}).json['totals']['planned'],4)

    def test_atomic_import_and_deduplication(self):
        (self.root/'IP24 - test.xlsx').touch()
        row={'order':'10','asset':'B','item':'I','planned':'2026-09-01','source':'a'}
        with patch('sap_plan.read_export',return_value=iter([row,row|{'source':'b'}])):
            result=import_exports(self.path,self.root)
        self.assertEqual(result['rows']['IP24'],1)
        self.assertEqual(result['duplicates'],1)
        with patch('sap_plan.read_export',return_value=iter([row,row|{'item_text':'conflict'}])):
            with self.assertRaises(ValueError): import_exports(self.path,self.root)
        with connect(self.path) as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM sap_rows WHERE kind='IP24'").fetchone()[0],1)

    def test_multiple_selections_and_month_drilldown(self):
        data=self.post('report',{'activity_type':['C6','T4'],'work_center':['ME','MM'],'month':'2026-09'}).json
        self.assertEqual(data['total'],3)
        self.assertEqual(data['totals']['planned'],4)
        self.assertEqual(data['by_plan'][0]['planned'],4)
        self.assertEqual(len(data['monthly']),2)
        self.assertEqual(self.post('report',{'activity_type':['T4'],'work_center':['MM']}).json['total'],0)
        self.assertEqual(self.post('report',{'activity_type':[],'work_center':[]}).json['total'],4)
        self.assertEqual(self.post('report',{'work_center':[42]}).status_code,400)

    def test_plan_only_join_is_not_ambiguous(self):
        with connect(self.path) as db:
            row={'order':'99','plan':'P','basic_start':'2026-09-01','system_status':'CONF'}
            db.execute('INSERT INTO sap_rows VALUES(?,?,?)',('IW39','99',json.dumps(row)))
        data=self.post('report',{}).json
        self.assertEqual(data['totals']['planned'],4)
        self.assertEqual(data['outside_master_orders'],2)

    def test_parent_116_includes_plan_on_descendant_055(self):
        parent='G-CSN-79-001-010-116'
        master={'plan':'LZCE1MA0245','item':'242548','asset':parent+'-055',
                'item_text':'REVISAR MOTOR BOMBA P11','source':'IP18'}
        order={'plan':master['plan'],'item':master['item'],'order':'80016227971',
               'asset':parent+'-055','basic_start':'2026-09-15','system_status':'CNPA','source':'IW39'}
        with connect(self.path) as db:
            for kind,key,row in [('IP18',master['item'],master),('IW39',order['order'],order)]:
                db.execute('INSERT INTO sap_rows VALUES(?,?,?)',(kind,key,json.dumps(row)))
        for selection in [{'assets':[parent]},{'search':parent}]:
            result=self.post('report',selection|{'date_from':'2026-09-01','date_to':'2026-09-30'}).json
            self.assertEqual(result['totals']['planned'],1)
            self.assertEqual(result['rows'][0]['order'],'80016227971')
            self.assertEqual(result['rows'][0]['asset_original'],parent+'-055')
            self.assertEqual(result['daily'][0]['planned'],1)

    def test_legacy_component_group_consolidates_without_duplicate(self):
        with connect(self.path) as db:
            db.execute('INSERT INTO sap_groups(name,assets) VALUES(?,?)',('Legado',json.dumps([
                'G-CSN-79-001-010-001-001','G-CSN-79-001-010-001-002'])))
        data=self.post('report',{'group':1}).json
        self.assertEqual(data['totals']['planned'],2)
        self.assertEqual(data['totals']['rate'],50)


if __name__=='__main__': unittest.main()
