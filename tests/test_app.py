import io
import sqlite3
import tempfile
import time
import unittest
from datetime import datetime
from pathlib import Path

import openpyxl

from app import create_app
from database import connect
from importer import import_file, classify_type, responsible_group

HEADERS = ['Equipto\n\n(PE)','Dt. Início','Dt. Fim','Tempo Parada Calculada (PE)',
           'Desc. Un. Resp. (PE)','Desc. Motivo Parada (PE)','Observação (PE)','Equipe (PE)','Turno (PE)']


def row(equipment='L1', start='2026-01-01 10:00', end='2026-01-01 10:10', minutes=10,
        responsible='ME-MECÂNICA', reason='Falha', observation='Sensor de saída'):
    return [equipment,datetime.fromisoformat(start),datetime.fromisoformat(end),minutes,responsible,reason,observation,'A',1]


def workbook(path, rows, headers=HEADERS):
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Paradas'
    ws.append(headers)
    for values in rows:
        ws.append(values)
        if isinstance(values[6],str) and values[6].startswith('='):
            ws.cell(ws.max_row,7).data_type='s'
    wb.save(path);wb.close()


class ApplicationTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.path=self.root/'test.sqlite3';self.xlsx=self.root/'2026.xlsx'
        self.app=create_app(self.path,self.root);self.app.config['TESTING']=True
        self.client=self.app.test_client();self.client.get('/')
        with self.client.session_transaction() as session:self.csrf=session['csrf']

    def tearDown(self):self.temp.cleanup()

    def post(self,path,body):return self.client.post(path,json=body,headers={'X-CSRF-Token':self.csrf})

    def load(self,rows,headers=HEADERS):
        workbook(self.xlsx,rows,headers);return import_file(self.path,self.xlsx)

    def summary(self,filters=None):
        response=self.post('/api/summary',filters or {});self.assertEqual(response.status_code,200,response.json)
        return response.json

    def test_consecutive_chain_different_reasons_and_equipment(self):
        self.load([row(),row(start='2026-01-01 10:11',end='2026-01-01 10:20',minutes=9,responsible='OP-OPERAÇÃO',reason='Ajuste'),
                   row(start='2026-01-01 10:21',end='2026-01-01 10:25',minutes=4),
                   row(start='2026-01-01 10:26:01',end='2026-01-01 10:30',minutes=4),row(equipment='L2')])
        totals=self.summary()['totals']
        self.assertEqual((totals['records'],totals['events'],totals['minutes']),(5,3,37))
        self.assertEqual(self.summary({'responsible':['Operação']})['totals']['events'],1)

    def test_rounding_and_overlapping_intervals(self):
        self.load([row(end='2026-01-01 10:10:59.999'),row(start='2026-01-01 10:12',end='2026-01-01 10:20'),
                   row(start='2026-01-01 10:15',end='2026-01-01 10:18')])
        totals=self.summary()['totals']
        self.assertEqual(totals['events'],1);self.assertEqual(totals['overlaps'],2)
        self.assertEqual(totals['minutes'],30)

    def test_offender_heatmap_distinct_events_and_filters(self):
        self.load([row(),row(start='2026-01-01 10:11',end='2026-01-01 10:20',minutes=9),
                   row(start='2026-02-01 10:00',end='2026-02-01 10:10',minutes=10)])
        offender=self.post('/api/offenders',{'name':'Sensor'}).json['id']
        records=self.post('/api/records',{}).json['rows']
        january=[r['id'] for r in records if r['start'].startswith('2026-01')]
        self.post('/api/classify',{'ids':january,'value':offender})
        data=self.summary()
        self.assertEqual(data['offender_months'],[
            {'period':'2026-01','key':str(offender),'minutes':19,'events':1},
            {'period':'2026-02','key':'none','minutes':10,'events':1}])
        filtered=self.summary({'offender':[offender],'month':[1]})
        self.assertEqual(filtered['offender_months'],data['offender_months'][:1])
        self.assertEqual(self.summary({'text':'inexistente'})['offender_months'],[])

    def test_annual_heatmap_counts_cross_month_event_once(self):
        self.load([row(start='2026-01-31 23:50',end='2026-02-01 00:00',minutes=10),
                   row(start='2026-02-01 00:01',end='2026-02-01 00:11',minutes=10),
                   row(start='2025-12-01 10:00',end='2025-12-01 10:10',minutes=10)])
        data=self.summary({'year':[2026]})
        self.assertEqual(data['offender_years'],[{'period':'2026','key':'none','minutes':20,'events':1}])
        self.assertEqual(sum(r['events'] for r in data['offender_months']),2)
        self.assertEqual(self.summary({'year':[2026],'month':[2]})['offender_years'],
                         [{'period':'2026','key':'none','minutes':10,'events':1}])

    def test_deduplication_and_reimport_preserve_classification(self):
        result=self.load([row(),row()]);self.assertEqual(result['duplicates'],1)
        offender=self.post('/api/offenders',{'name':'Sensor'}).json['id']
        record=self.post('/api/records',{}).json['rows'][0]['id']
        self.assertEqual(self.post('/api/classify',{'ids':[record],'value':offender}).status_code,200)
        self.assertEqual(import_file(self.path,self.xlsx)['status'],'unchanged')
        self.load([row(),row(equipment='L2')])
        self.assertEqual(self.summary({'offender':[offender]})['totals']['records'],1)
        copy=self.root/'copia.xlsx';workbook(copy,[row()]);import_file(self.path,copy)
        self.assertEqual(self.summary()['totals']['records'],2)
        meta=self.client.get('/api/meta').json
        ids=[f['id'] for f in meta['files']]
        self.assertEqual(self.summary({'files':ids})['totals']['records'],2)

    def test_replaced_file_hides_old_records_but_can_recover_tags(self):
        self.load([row()]);offender=self.post('/api/offenders',{'name':'Sensor'}).json['id']
        self.post('/api/classify',{'all_filtered':True,'filters':{},'value':offender})
        self.load([row(equipment='L2')]);self.assertEqual(self.summary({'offender':[offender]})['totals']['records'],0)
        self.load([row()]);self.assertEqual(self.summary({'offender':[offender]})['totals']['records'],1)

    def test_alternate_headers_and_validation(self):
        headers=['Equipamento','Dt. Início','Dt. Fim','Tempo Parada Calculada','Desc. Un. Resp.','Desc. Motivo Parada','Observação','Equipe','Turno']
        result=self.load([row(),row(minutes=-1),row(start='2026-01-01 12:00',end='2026-01-01 10:00')],headers)
        self.assertEqual(result['invalid'],2);self.assertEqual(self.summary()['totals']['records'],1)
        workbook(self.xlsx,[row(minutes=-1)],headers)
        with self.assertRaises(ValueError):import_file(self.path,self.xlsx)
        self.assertEqual(self.summary()['totals']['records'],1)

    def test_filters_accents_dates_and_source_coverage(self):
        self.load([row(),row(start='2026-02-10 10:00',end='2026-02-10 10:20',minutes=20,observation='Tesoura'),
                   row(start='2025-12-31 23:55',end='2026-01-01 00:05')])
        self.assertEqual(self.summary({'text':'SENSOR SAIDA'})['totals']['records'],2)
        self.assertEqual(self.summary({'year':['2026'],'month':['2'],'day':['10'],'min_minutes':15})['totals']['records'],1)
        self.assertEqual(self.summary({'date_from':'2026-01-01','date_to':'2026-01-31'})['totals']['records'],1)
        empty=self.summary({'year':['2026'],'text':'inexistente'})
        self.assertEqual(len(empty['coverage']),2);self.assertEqual(len(empty['periods']),12)
        self.assertEqual(self.post('/api/summary',{'month':[13]}).status_code,400)
        self.assertEqual(self.post('/api/summary',{'min_minutes':30,'max_minutes':10}).status_code,400)

    def test_bulk_classification_and_manual_type(self):
        self.load([row(),row(equipment='L2')])
        offender=self.post('/api/offenders',{'name':'Sensor'}).json['id']
        response=self.post('/api/classify',{'all_filtered':True,'filters':{'equipment':['L1']},'expected_count':1,'value':offender})
        self.assertEqual(response.json['updated'],1)
        self.assertEqual(self.summary({'offender':['none']})['totals']['records'],1)
        result=self.post('/api/classify',{'all_filtered':True,'filters':{},'expected_count':1,'value':offender})
        self.assertEqual(result.status_code,400)
        self.post('/api/classify',{'all_filtered':True,'filters':{},'field':'type_override','value':'Reparo Geral'})
        self.assertEqual(self.summary({'type':['Reparo Geral']})['totals']['records'],2)
        self.post('/api/classify',{'all_filtered':True,'filters':{},'field':'type_override','value':None})
        self.assertEqual(self.summary({'type':['Reparo Geral']})['totals']['records'],0)

    def test_event_frequency_not_launch_count(self):
        self.load([row(),row(start='2026-01-01 10:11',end='2026-01-01 10:20',minutes=9)])
        offender=self.post('/api/offenders',{'name':'Sensor'}).json['id']
        self.post('/api/classify',{'all_filtered':True,'filters':{},'value':offender})
        group=self.summary()['groups']['offender'][0]
        self.assertEqual((group['events'],group['records'],group['minutes']),(1,2,19))
        record=self.post('/api/records',{}).json['rows'][0]['id']
        detail=self.client.get('/api/records/'+record).json
        self.assertEqual(len(detail['event_rows']),2);self.assertEqual(detail['sources'][0]['sheet'],'Paradas')
        self.post('/api/settings',{'tolerance':0})
        self.assertEqual(self.summary()['totals']['events'],2)

    def test_export_backup_and_csrf(self):
        self.load([row(observation='=HYPERLINK("example")')])
        self.assertEqual(self.client.post('/api/offenders',json={'name':'x'}).status_code,403)
        exported=self.post('/api/export',{}).data.decode('utf-8-sig')
        self.assertIn("'=HYPERLINK",exported);self.assertIn('linha 2',exported)
        backup=self.client.get('/api/backup');self.assertEqual(backup.status_code,200)
        path=self.root/'backup.sqlite3';path.write_bytes(backup.data)
        with connect(path) as db:self.assertEqual(db.execute('SELECT COUNT(*) FROM records').fetchone()[0],1)

    def test_type_and_responsible_rules(self):
        self.assertEqual(classify_type('PR-MP',''),'Manutenção Preventiva')
        self.assertEqual(classify_type('','Reparo geral da linha'),'Reparo Geral')
        self.assertEqual(classify_type('','tempo de parada'),'Demais paradas')
        self.assertEqual(responsible_group('PA-PROCESSO ANTERIOR'),'Externa')
        self.assertEqual(responsible_group('PR-PROGRAMAÇÃO'),'Programação')
        self.assertEqual(responsible_group('EL-ELÉTRICA'),'Elétrica')
        self.assertEqual(responsible_group('ME-MANUTENÇÃO ELÉTRICA'),'Elétrica')
        self.assertEqual(responsible_group('MM-MANUTENÇÃO MECÂNICA'),'Mecânica')

    def wait_import(self):
        deadline=time.monotonic()+10
        while time.monotonic()<deadline:
            status=self.client.get('/api/import/status').json
            if not status['running']:return status
            time.sleep(.05)
        self.fail('Importação não terminou em 10 segundos.')

    def test_local_import_job_and_path_validation(self):
        workbook(self.xlsx,[row()])
        self.assertEqual(self.post('/api/import/local',{'names':['../outro.xlsx']}).status_code,400)
        self.assertEqual(self.post('/api/import/local',{'names':['2026.xlsx']}).status_code,202)
        status=self.wait_import();self.assertEqual(status['results'][0]['status'],'imported')
        self.assertEqual(self.summary()['totals']['records'],1)

    def test_uploaded_workbook_job(self):
        workbook(self.xlsx,[row()])
        response=self.client.post('/api/import/upload',data={'files':(io.BytesIO(self.xlsx.read_bytes()),'enviado.xlsx')},
                                  headers={'X-CSRF-Token':self.csrf},content_type='multipart/form-data')
        self.assertEqual(response.status_code,202)
        self.assertEqual(self.wait_import()['results'][0]['status'],'imported')
        self.assertEqual(self.client.get('/api/meta').json['files'][0]['name'],'enviado.xlsx')


if __name__=='__main__':unittest.main()
